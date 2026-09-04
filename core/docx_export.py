import os
import datetime
from docx import Document

DOCX_OUTPUT_DIR = os.path.join(os.path.expanduser("~"), "Documents", "GreekG_Exports")


def save_clipboard_to_docx(content: str) -> tuple[bool, str]:
    try:
        if not content or not content.strip():
            return False, "Clipboard is empty"

        os.makedirs(DOCX_OUTPUT_DIR, exist_ok=True)
        today_str = datetime.date.today().isoformat()
        doc_path = os.path.join(DOCX_OUTPUT_DIR, f"GreekG_{today_str}.docx")

        if os.path.exists(doc_path):
            doc = Document(doc_path)
        else:
            doc = Document()
            doc.add_heading(f"GreekG Session Log — {today_str}", level=1)

        timestamp = datetime.datetime.now().strftime("%H:%M:%S")
        doc.add_heading(f"Entry — {timestamp}", level=2)
        doc.add_paragraph(content.strip())
        doc.save(doc_path)
        return True, f"Saved to GreekG_{today_str}.docx"
    except Exception as e:
        return False, f"Export error: {e}"
